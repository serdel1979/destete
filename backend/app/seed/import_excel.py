"""Migra los datos reales de los dos Excel del cliente a Postgres.

Fuente 1: "Seguimiento Terneros.xls" -> lotes históricos, animales y pesajes.
Fuente 2: "2026_Control_Destete_GENITU...xlsx" -> lote de planificación 2026
          (Caravana Celeste) con su plan de alimentación/costos.

Se ejecuta una sola vez (es idempotente: si ya hay lotes cargados, no hace nada).
"""
from datetime import date, datetime
from pathlib import Path

import openpyxl
import xlrd
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models.alimentacion import PlanAlimentacion
from app.models.animal import Animal, EstadoAnimal, SexoAnimal
from app.models.lote import Lote
from app.models.pesaje import Pesaje

DATA_DIR = Path(__file__).resolve().parents[2] / "data" / "raw"
SEGUIMIENTO_PATH = DATA_DIR / "Seguimiento Terneros.xls"
GENITU_PATH = DATA_DIR / "2026_Control_Destete_GENITU_Lote 1 - Caravana_Celeste.xlsx"

# Fechas de pesaje conocidas (fila "Fecha de Pesaje" de la planilla) para las
# columnas de PESO 1 a 5. Las columnas 6 y 7 no tienen fecha asociada en el
# original y se importan como observación, no como pesaje formal.
COLS_PESO_CON_FECHA = [8, 9, 10, 11, 12]
COLS_PESO_SIN_FECHA = [13, 14]


def _get_or_create_lote(session: Session, nombre: str, **kwargs) -> Lote:
    lote = session.execute(select(Lote).where(Lote.nombre == nombre)).scalar_one_or_none()
    if lote:
        return lote
    lote = Lote(nombre=nombre, **kwargs)
    session.add(lote)
    session.flush()
    return lote


def _xldate(wb: xlrd.Book, value) -> date | None:
    if value in (None, "", 0):
        return None
    try:
        return xlrd.xldate_as_datetime(float(value), wb.datemode).date()
    except (ValueError, TypeError):
        return None


def _unique_caravana(session: Session, base: str, seen: dict[str, int]) -> tuple[str, bool]:
    """Genera una caravana única. Devuelve (caravana, es_provisoria)."""
    existing = session.execute(select(Animal).where(Animal.caravana == base)).scalar_one_or_none()
    if not existing and base not in seen:
        seen[base] = 0
        return base, False
    seen[base] = seen.get(base, 0) + 1
    return f"{base}-DUP{seen[base]}", True


def import_seguimiento_terneros(session: Session) -> None:
    wb = xlrd.open_workbook(str(SEGUIMIENTO_PATH), formatting_info=False)
    ws = wb.sheet_by_name("Hoja1")

    lotes_cache: dict[str, Lote] = {}
    sc_counters: dict[str, int] = {}
    caravana_seen: dict[str, int] = {}

    total_animales = 0
    total_pesajes = 0

    for r in range(2, 184):  # filas 2..183: registro individual de terneros
        row = ws.row_values(r)
        if not any(row):
            continue

        raw_caravana = row[0]
        sexo_raw = str(row[1]).strip().upper() if row[1] else ""
        if sexo_raw not in ("M", "H"):
            continue  # fila sin datos utilizables

        caracteristicas = str(row[2]).strip() if row[2] else None
        fecha_nacimiento = _xldate(wb, row[3])
        fecha_destete = _xldate(wb, row[4])
        madre = str(row[5]).strip() if row[5] else None
        padre = str(row[6]).strip() if row[6] else None
        lote_nombre = str(row[7]).strip() if row[7] else None

        lote = None
        if lote_nombre:
            lote = lotes_cache.get(lote_nombre)
            if not lote:
                lote = _get_or_create_lote(
                    session,
                    lote_nombre,
                    estado="activo",
                    observaciones="Migrado desde Seguimiento Terneros.xls",
                )
                lotes_cache[lote_nombre] = lote

        caravana_provisoria = False
        if raw_caravana in (None, ""):
            base = f"SIN-CARAVANA-F{r}"
            caravana_provisoria = True
        elif isinstance(raw_caravana, float) and raw_caravana.is_integer():
            base = str(int(raw_caravana))
        else:
            texto = str(raw_caravana).strip()
            if texto.upper().startswith("S/C"):
                sc_counters[sexo_raw] = sc_counters.get(sexo_raw, 0) + 1
                base = f"SC-{sexo_raw}-{sc_counters[sexo_raw]:03d}"
                caravana_provisoria = True
            else:
                base = texto

        caravana, dup = _unique_caravana(session, base, caravana_seen)
        caravana_provisoria = caravana_provisoria or dup

        observaciones_extra = []
        if raw_caravana and str(raw_caravana).strip().upper().startswith("S/C"):
            observaciones_extra.append(f"Caravana original en planilla: '{raw_caravana}'")
        for c in COLS_PESO_SIN_FECHA:
            if len(row) > c and row[c] not in (None, "", 0):
                observaciones_extra.append(
                    f"Peso adicional sin fecha registrada en planilla original: {row[c]} kg"
                )
        if len(row) > 17 and any(v not in (None, "") for v in row[15:18]):
            observaciones_extra.append(f"Nota planilla original: {row[15:18]}")

        animal = Animal(
            caravana=caravana,
            caravana_provisoria=caravana_provisoria,
            sexo=SexoAnimal.macho if sexo_raw == "M" else SexoAnimal.hembra,
            caracteristicas=caracteristicas,
            fecha_nacimiento=fecha_nacimiento,
            fecha_destete=fecha_destete,
            madre=madre,
            padre=padre,
            lote_id=lote.id if lote else None,
            estado=EstadoAnimal.activo,
            observaciones="; ".join(observaciones_extra) or None,
        )
        session.add(animal)
        session.flush()
        total_animales += 1

        for col in COLS_PESO_CON_FECHA:
            if col >= len(row):
                continue
            peso = row[col]
            if peso in (None, "", 0):
                continue
            fecha_pesaje = _xldate(wb, ws.cell(196, col).value)
            if not fecha_pesaje:
                continue
            session.add(Pesaje(animal_id=animal.id, fecha=fecha_pesaje, peso_kg=float(peso)))
            total_pesajes += 1

    session.commit()
    print(f"Seguimiento Terneros: {total_animales} animales, {total_pesajes} pesajes importados.")


def import_control_destete(session: Session) -> None:
    wb = openpyxl.load_workbook(str(GENITU_PATH), data_only=True)

    variables = {}
    for row in wb["VARIABLES MODELO"].iter_rows(values_only=True):
        if row and row[0] and row[1] is not None:
            variables[str(row[0]).strip()] = row[1]

    fecha_inicio_raw = variables.get("Fecha inicio destete")
    fecha_inicio = fecha_inicio_raw.date() if isinstance(fecha_inicio_raw, datetime) else None
    peso_inicial = float(variables.get("Peso inicial", 0) or 0)

    lote = _get_or_create_lote(
        session,
        "Lote 1 - Caravana Celeste 2026",
        caravana_color="Celeste",
        fecha_inicio=fecha_inicio,
        peso_inicial_promedio=peso_inicial,
        estado="activo",
        observaciones=(
            f"Migrado desde planilla de planificación GENITU. "
            f"Cantidad de terneros planificada: {int(variables.get('Cantidad de terneros', 0))}."
        ),
    )

    precio_hiper = float(variables.get("Precio Hiper-Precoz 23% p/kg", 0) or 0)
    precio_precoz = float(variables.get("Precio Precoz 19% p/kg", 0) or 0)

    # Segmentos derivados de VARIABLES MODELO (% consumo) + MODELO 150D
    # (fase/ADPV promedio real por tramo). Ver docs/propuesta.md para el
    # detalle de por qué se simplifica a ADPV constante por tramo.
    segmentos = [
        ("Hiper-Precoz 23%", 1, 5, 0.0200, -0.100, precio_hiper),
        ("Hiper-Precoz 23%", 6, 15, 0.0350, 0.070, precio_hiper),
        ("Precoz 19%", 16, 50, 0.0320, 0.381, precio_precoz),
        ("Precoz 19%", 51, 80, 0.0300, 0.572, precio_precoz),
        ("Precoz 19%", 81, 120, 0.0280, 0.654, precio_precoz),
        ("Precoz 19%", 121, 150, 0.0250, 0.700, precio_precoz),
    ]

    existing = session.execute(
        select(PlanAlimentacion).where(PlanAlimentacion.lote_id == lote.id)
    ).scalars().all()
    if not existing:
        for i, (fase, desde, hasta, pct, adpv, costo) in enumerate(segmentos, start=1):
            session.add(
                PlanAlimentacion(
                    lote_id=lote.id,
                    orden=i,
                    fase=fase,
                    dia_desde=desde,
                    dia_hasta=hasta,
                    pct_consumo_pv=pct,
                    adpv_esperado_kg=adpv,
                    costo_kg=costo,
                )
            )
        print(f"Plan de alimentación cargado para '{lote.nombre}' ({len(segmentos)} etapas).")

    session.commit()


def run() -> None:
    engine = create_engine(settings.DATABASE_URL_SYNC)
    with Session(engine) as session:
        ya_migrado = session.execute(select(Lote)).scalars().first()
        if ya_migrado:
            print("Ya existen lotes en la base de datos, se omite la migración.")
            return
        import_seguimiento_terneros(session)
        import_control_destete(session)


if __name__ == "__main__":
    run()
